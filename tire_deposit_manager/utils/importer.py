#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""
Moduł do importu danych z plików Excel i CSV do bazy danych.
"""

import os
import logging
import csv
from datetime import datetime

# Logger
logger = logging.getLogger("TireDepositManager")

def import_data_from_excel(conn, file_path, data_type, sheet_name=None):
    """
    Importuje dane z pliku Excel do bazy danych.
    
    Args:
        conn: Połączenie z bazą danych
        file_path (str): Ścieżka do pliku źródłowego
        data_type (str): Typ danych (np. 'clients', 'deposits', 'inventory')
        sheet_name (str, optional): Nazwa arkusza z danymi. Jeśli None, używa pierwszego.
        
    Returns:
        int: Liczba zaimportowanych rekordów
    """
    try:
        # Upewnijmy się, że mamy dostępne biblioteki
        try:
            import openpyxl
        except ImportError:
            logger.error("Brak biblioteki openpyxl. Zainstaluj ją: pip install openpyxl")
            raise ImportError("Brak biblioteki openpyxl. Zainstaluj ją: pip install openpyxl")
        
        # Wczytanie pliku Excel
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        
        # Wybór arkusza
        if sheet_name and sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active
        
        # Pobranie wierszy z arkusza
        rows = list(worksheet.rows)
        
        # Jeśli nie ma danych, zwróć 0
        if not rows:
            logger.warning(f"Brak danych w arkuszu: {worksheet.title}")
            return 0
        
        # Pobranie nagłówków (pierwszy wiersz)
        headers = [cell.value for cell in rows[0]]
        
        # Przygotowanie danych do importu
        data_rows = []
        for row in rows[1:]:  # Pomijamy wiersz nagłówkowy
            row_data = {}
            for i, cell in enumerate(row):
                if i < len(headers) and headers[i] is not None:
                    # Formatowanie komórek
                    if cell.value is not None:
                        row_data[headers[i]] = cell.value
            
            # Dodaj wiersz tylko jeśli nie jest pusty
            if row_data:
                data_rows.append(row_data)
        
        # Importowanie danych do odpowiedniej tabeli
        count = 0
        
        if data_type == "clients":
            count = import_clients(conn, data_rows, headers)
        elif data_type == "deposits":
            count = import_deposits(conn, data_rows, headers)
        elif data_type == "inventory":
            count = import_inventory(conn, data_rows, headers)
        elif data_type == "parts":
            count = import_parts(conn, data_rows, headers)
        elif data_type == "appointments":
            count = import_appointments(conn, data_rows, headers)
        else:
            logger.error(f"Nieznany typ danych: {data_type}")
            return 0
        
        logger.info(f"Zaimportowano {count} rekordów typu {data_type} z pliku Excel: {file_path}")
        return count
    
    except Exception as e:
        logger.error(f"Błąd podczas importu z Excel: {e}")
        raise


def import_data_from_csv(conn, file_path, data_type):
    """
    Importuje dane z pliku CSV do bazy danych.
    
    Args:
        conn: Połączenie z bazą danych
        file_path (str): Ścieżka do pliku źródłowego
        data_type (str): Typ danych (np. 'clients', 'deposits', 'inventory')
        
    Returns:
        int: Liczba zaimportowanych rekordów
    """
    try:
        # Wykryj separator (średnik lub przecinek)
        with open(file_path, 'r', encoding='utf-8-sig') as f:
            first_line = f.readline()
            if ';' in first_line:
                delimiter = ';'
            else:
                delimiter = ','
        
        # Wczytanie pliku CSV
        with open(file_path, 'r', encoding='utf-8-sig') as csv_file:
            csv_reader = csv.DictReader(csv_file, delimiter=delimiter)
            
            # Pobranie nagłówków
            headers = csv_reader.fieldnames
            
            # Przygotowanie danych do importu
            data_rows = []
            for row in csv_reader:
                # Dodaj wiersz tylko jeśli nie jest pusty
                if any(row.values()):
                    data_rows.append(row)
        
        # Importowanie danych do odpowiedniej tabeli
        count = 0
        
        if data_type == "clients":
            count = import_clients(conn, data_rows, headers)
        elif data_type == "deposits":
            count = import_deposits(conn, data_rows, headers)
        elif data_type == "inventory":
            count = import_inventory(conn, data_rows, headers)
        elif data_type == "parts":
            count = import_parts(conn, data_rows, headers)
        elif data_type == "appointments":
            count = import_appointments(conn, data_rows, headers)
        else:
            logger.error(f"Nieznany typ danych: {data_type}")
            return 0
        
        logger.info(f"Zaimportowano {count} rekordów typu {data_type} z pliku CSV: {file_path}")
        return count
    
    except Exception as e:
        logger.error(f"Błąd podczas importu z CSV: {e}")
        raise


def import_clients(conn, data_rows, headers):
    """
    Importuje dane klientów do bazy danych.
    
    Args:
        conn: Połączenie z bazą danych
        data_rows (list): Lista słowników z danymi klientów
        headers (list): Lista nagłówków
        
    Returns:
        int: Liczba zaimportowanych rekordów
    """
    try:
        cursor = conn.cursor()
        count = 0
        
        # Mapowanie nagłówków do kolumn w bazie danych
        header_mapping = {
            'Nazwa klienta': 'name',
            'Telefon': 'phone_number',
            'E-mail': 'email',
            'Informacje dodatkowe': 'additional_info',
            'Rabat (%)': 'discount',
            'Kod kreskowy': 'barcode',
            # Angielskie wersje nagłówków
            'Name': 'name',
            'Phone': 'phone_number',
            'Email': 'email',
            'Additional Info': 'additional_info',
            'Discount (%)': 'discount',
            'Barcode': 'barcode'
        }
        
        # Dodaj klientów do bazy danych
        for row in data_rows:
            # Przygotuj dane klienta w formacie odpowiednim dla bazy danych
            client_data = {}
            for header in row.keys():
                if header in header_mapping and row[header] is not None:
                    client_data[header_mapping[header]] = row[header]
            
            # Jeśli brak wymaganych danych, pomiń wiersz
            if 'name' not in client_data or not client_data['name']:
                continue
            
            # Sprawdź, czy klient już istnieje (po nazwie i telefonie)
            cursor.execute(
                "SELECT id FROM clients WHERE name = ? AND (phone_number = ? OR phone_number IS NULL)",
                (client_data.get('name'), client_data.get('phone_number', None))
            )
            existing_client = cursor.fetchone()
            
            if existing_client:
                # Aktualizuj istniejącego klienta
                update_columns = []
                update_values = []
                
                for column, value in client_data.items():
                    if column != 'name':  # Nie aktualizujemy nazwy
                        update_columns.append(f"{column} = ?")
                        update_values.append(value)
                
                if update_columns:
                    update_query = f"UPDATE clients SET {', '.join(update_columns)} WHERE id = ?"
                    update_values.append(existing_client[0])
                    cursor.execute(update_query, update_values)
                    count += 1
            else:
                # Dodaj nowego klienta
                columns = list(client_data.keys())
                placeholders = ['?'] * len(columns)
                values = [client_data[column] for column in columns]
                
                insert_query = f"INSERT INTO clients ({', '.join(columns)}) VALUES ({', '.join(placeholders)})"
                cursor.execute(insert_query, values)
                count += 1
        
        # Zatwierdzenie zmian
        conn.commit()
        
        return count
    
    except Exception as e:
        conn.rollback()
        logger.error(f"Błąd podczas importu klientów: {e}")
        raise


def import_deposits(conn, data_rows, headers):
    """
    Importuje dane depozytów do bazy danych.
    
    Args:
        conn: Połączenie z bazą danych
        data_rows (list): Lista słowników z danymi depozytów
        headers (list): Lista nagłówków
        
    Returns:
        int: Liczba zaimportowanych rekordów
    """
    try:
        cursor = conn.cursor()
        count = 0
        
        # Mapowanie nagłówków do kolumn w bazie danych
        header_mapping = {
            'Klient': 'client_name',
            'Model auta': 'car_model',
            'Nr rejestracyjny': 'registration_number',
            'Marka opon': 'tire_brand',
            'Rozmiar opon': 'tire_size',
            'Ilość': 'quantity',
            'Lokalizacja': 'location',
            'Data depozytu': 'deposit_date',
            'Oczekiwany zwrot': 'expected_return_date',
            'Status': 'status',
            'Sezon': 'season',
            'Cena (PLN)': 'price',
            # Angielskie wersje nagłówków
            'Client': 'client_name',
            'Car Model': 'car_model',
            'Registration Number': 'registration_number',
            'Tire Brand': 'tire_brand',
            'Tire Size': 'tire_size',
            'Quantity': 'quantity',
            'Location': 'location',
            'Deposit Date': 'deposit_date',
            'Expected Return': 'expected_return_date',
            'Status': 'status',
            'Season': 'season',
            'Price (PLN)': 'price'
        }
        
        # Dodaj depozyty do bazy danych
        for row in data_rows:
            # Przygotuj dane depozytu w formacie odpowiednim dla bazy danych
            deposit_data = {}
            for header in row.keys():
                if header in header_mapping and row[header] is not None:
                    field_name = header_mapping[header]
                    value = row[header]
                    
                    # Konwersja formatu daty
                    if field_name in ['deposit_date', 'expected_return_date']:
                        try:
                            if isinstance(value, str):
                                # Sprawdź różne formaty daty
                                if '.' in value:  # Format dd.mm.yyyy
                                    date_obj = datetime.strptime(value, "%d.%m.%Y")
                                elif '-' in value and len(value) >= 10:  # Format yyyy-mm-dd
                                    date_obj = datetime.strptime(value[:10], "%Y-%m-%d")
                                else:
                                    # Spróbuj domyślny format Excela
                                    date_obj = datetime.strptime(value, "%Y-%m-%d")
                                
                                value = date_obj.strftime("%Y-%m-%d")
                            elif hasattr(value, 'strftime'):  # Obiekt daty
                                value = value.strftime("%Y-%m-%d")
                        except (ValueError, TypeError):
                            # Jeśli nie udało się przekonwertować, użyj oryginalnej wartości
                            pass
                    
                    # Konwersja ceny
                    elif field_name == 'price':
                        try:
                            if isinstance(value, str):
                                # Usuń znaki PLN i zamień przecinki na kropki
                                value = value.replace(' PLN', '').replace(',', '.')
                            value = float(value)
                        except (ValueError, TypeError):
                            # Jeśli nie udało się przekonwertować, ustaw domyślną cenę
                            value = 0.0
                    
                    # Konwersja ilości
                    elif field_name == 'quantity':
                        try:
                            value = int(value)
                        except (ValueError, TypeError):
                            # Jeśli nie udało się przekonwertować, ustaw domyślną ilość
                            value = 0
                    
                    deposit_data[field_name] = value
            
            # Jeśli brak wymaganych danych, pomiń wiersz
            if 'registration_number' not in deposit_data or not deposit_data['registration_number']:
                continue
            
            # Znajdź ID klienta na podstawie nazwy
            client_id = None
            if 'client_name' in deposit_data:
                cursor.execute("SELECT id FROM clients WHERE name = ?", (deposit_data['client_name'],))
                client_result = cursor.fetchone()
                
                if client_result:
                    client_id = client_result[0]
                else:
                    # Jeśli klient nie istnieje, dodaj go
                    cursor.execute("INSERT INTO clients (name) VALUES (?)", (deposit_data['client_name'],))
                    client_id = cursor.lastrowid
            
            # Usuń nazwę klienta z danych depozytu
            if 'client_name' in deposit_data:
                del deposit_data['client_name']
            
            # Dodaj ID klienta do danych depozytu
            if client_id:
                deposit_data['client_id'] = client_id
            
            # Sprawdź, czy depozyt już istnieje (po numerze rejestracyjnym i marce opon)
            cursor.execute(
                "SELECT id FROM deposits WHERE registration_number = ? AND tire_brand = ?",
                (deposit_data.get('registration_number'), deposit_data.get('tire_brand', ''))
            )
            existing_deposit = cursor.fetchone()
            
            if existing_deposit:
                # Aktualizuj istniejący depozyt
                update_columns = []
                update_values = []
                
                for column, value in deposit_data.items():
                    if column not in ['registration_number', 'tire_brand']:  # Nie aktualizujemy kluczy
                        update_columns.append(f"{column} = ?")
                        update_values.append(value)
                
                if update_columns:
                    update_query = f"UPDATE deposits SET {', '.join(update_columns)} WHERE id = ?"
                    update_values.append(existing_deposit[0])
                    cursor.execute(update_query, update_values)
                    count += 1
            else:
                # Dodaj nowy depozyt
                columns = list(deposit_data.keys())
                placeholders = ['?'] * len(columns)
                values = [deposit_data[column] for column in columns]
                
                # Dodaj domyślny status, jeśli nie podano
                if 'status' not in deposit_data:
                    columns.append('status')
                    placeholders.append('?')
                    values.append('Aktywny')
                
                insert_query = f"INSERT INTO deposits ({', '.join(columns)}) VALUES ({', '.join(placeholders)})"
                cursor.execute(insert_query, values)
                count += 1
        
        # Zatwierdzenie zmian
        conn.commit()
        
        return count
    
    except Exception as e:
        conn.rollback()
        logger.error(f"Błąd podczas importu depozytów: {e}")
        raise


def import_inventory(conn, data_rows, headers):
    """
    Importuje dane inwentarza opon do bazy danych.
    
    Args:
        conn: Połączenie z bazą danych
        data_rows (list): Lista słowników z danymi inwentarza
        headers (list): Lista nagłówków
        
    Returns:
        int: Liczba zaimportowanych rekordów
    """
    try:
        cursor = conn.cursor()
        count = 0
        
        # Mapowanie nagłówków do kolumn w bazie danych
        header_mapping = {
            'Marka i model': 'brand_model',
            'Rozmiar': 'size',
            'Ilość': 'quantity',
            'Cena (PLN)': 'price',
            'DOT': 'dot',
            'Sezon': 'season_type',
            'Uwagi': 'notes',
            # Angielskie wersje nagłówków
            'Brand and Model': 'brand_model',
            'Size': 'size',
            'Quantity': 'quantity',
            'Price (PLN)': 'price',
            'DOT': 'dot',
            'Season': 'season_type',
            'Notes': 'notes'
        }
        
        # Dodaj inwentarz do bazy danych
        for row in data_rows:
            # Przygotuj dane inwentarza w formacie odpowiednim dla bazy danych
            inventory_data = {}
            for header in row.keys():
                if header in header_mapping and row[header] is not None:
                    field_name = header_mapping[header]
                    value = row[header]
                    
                    # Konwersja ceny
                    if field_name == 'price':
                        try:
                            if isinstance(value, str):
                                # Usuń znaki PLN i zamień przecinki na kropki
                                value = value.replace(' PLN', '').replace(',', '.')
                            value = float(value)
                        except (ValueError, TypeError):
                            # Jeśli nie udało się przekonwertować, ustaw domyślną cenę
                            value = 0.0
                    
                    # Konwersja ilości
                    elif field_name == 'quantity':
                        try:
                            value = int(value)
                        except (ValueError, TypeError):
                            # Jeśli nie udało się przekonwertować, ustaw domyślną ilość
                            value = 0
                    
                    inventory_data[field_name] = value
            
            # Jeśli brak wymaganych danych, pomiń wiersz
            if 'brand_model' not in inventory_data or not inventory_data['brand_model'] or 'size' not in inventory_data or not inventory_data['size']:
                continue
            
            # Sprawdź, czy pozycja inwentarza już istnieje (po marce/modelu i rozmiarze)
            cursor.execute(
                "SELECT id FROM inventory WHERE brand_model = ? AND size = ?",
                (inventory_data.get('brand_model'), inventory_data.get('size'))
            )
            existing_item = cursor.fetchone()
            
            if existing_item:
                # Aktualizuj istniejącą pozycję
                update_columns = []
                update_values = []
                
                for column, value in inventory_data.items():
                    if column not in ['brand_model', 'size']:  # Nie aktualizujemy kluczy
                        update_columns.append(f"{column} = ?")
                        update_values.append(value)
                
                if update_columns:
                    update_query = f"UPDATE inventory SET {', '.join(update_columns)} WHERE id = ?"
                    update_values.append(existing_item[0])
                    cursor.execute(update_query, update_values)
                    count += 1
            else:
                # Dodaj nową pozycję
                columns = list(inventory_data.keys())
                placeholders = ['?'] * len(columns)
                values = [inventory_data[column] for column in columns]
                
                insert_query = f"INSERT INTO inventory ({', '.join(columns)}) VALUES ({', '.join(placeholders)})"
                cursor.execute(insert_query, values)
                count += 1
        
        # Zatwierdzenie zmian
        conn.commit()
        
        return count
    
    except Exception as e:
        conn.rollback()
        logger.error(f"Błąd podczas importu inwentarza: {e}")
        raise


def import_parts(conn, data_rows, headers):
    """
    Importuje dane części i akcesoriów do bazy danych.
    
    Args:
        conn: Połączenie z bazą danych
        data_rows (list): Lista słowników z danymi części
        headers (list): Lista nagłówków
        
    Returns:
        int: Liczba zaimportowanych rekordów
    """
    try:
        cursor = conn.cursor()
        count = 0
        
        # Mapowanie nagłówków do kolumn w bazie danych
        header_mapping = {
            'Nazwa': 'name',
            'Nr katalogowy': 'catalog_number',
            'Kategoria': 'category',
            'Producent': 'manufacturer',
            'Ilość': 'quantity',
            'Cena (PLN)': 'price',
            'Lokalizacja': 'location',
            'Opis': 'description',
            'Kod kreskowy': 'barcode',
            'Minimalna ilość': 'minimum_quantity',
            # Angielskie wersje nagłówków
            'Name': 'name',
            'Catalog Number': 'catalog_number',
            'Category': 'category',
            'Manufacturer': 'manufacturer',
            'Quantity': 'quantity',
            'Price (PLN)': 'price',
            'Location': 'location',
            'Description': 'description',
            'Barcode': 'barcode',
            'Minimum Quantity': 'minimum_quantity'
        }
        
        # Dodaj części do bazy danych
        for row in data_rows:
            # Przygotuj dane części w formacie odpowiednim dla bazy danych
            part_data = {}
            for header in row.keys():
                if header in header_mapping and row[header] is not None:
                    field_name = header_mapping[header]
                    value = row[header]
                    
                    # Konwersja ceny
                    if field_name == 'price':
                        try:
                            if isinstance(value, str):
                                # Usuń znaki PLN i zamień przecinki na kropki
                                value = value.replace(' PLN', '').replace(',', '.')
                            value = float(value)
                        except (ValueError, TypeError):
                            # Jeśli nie udało się przekonwertować, ustaw domyślną cenę
                            value = 0.0
                    
                    # Konwersja ilości
                    elif field_name in ['quantity', 'minimum_quantity']:
                        try:
                            value = int(value)
                        except (ValueError, TypeError):
                            # Jeśli nie udało się przekonwertować, ustaw domyślną ilość
                            value = 0
                    
                    part_data[field_name] = value
            
            # Jeśli brak wymaganych danych, pomiń wiersz
            if 'name' not in part_data or not part_data['name']:
                continue
            
            # Sprawdź, czy część już istnieje (po nazwie i numerze katalogowym)
            # Jeśli numer katalogowy nie jest dostępny, sprawdź tylko po nazwie
            if 'catalog_number' in part_data and part_data['catalog_number']:
                cursor.execute(
                    "SELECT id FROM parts WHERE name = ? AND catalog_number = ?",
                    (part_data.get('name'), part_data.get('catalog_number'))
                )
            else:
                cursor.execute(
                    "SELECT id FROM parts WHERE name = ?",
                    (part_data.get('name'),)
                )
            
            existing_part = cursor.fetchone()
            
            if existing_part:
                # Aktualizuj istniejącą część
                update_columns = []
                update_values = []
                
                for column, value in part_data.items():
                    if column != 'name':  # Nie aktualizujemy nazwy
                        update_columns.append(f"{column} = ?")
                        update_values.append(value)
                
                if update_columns:
                    update_query = f"UPDATE parts SET {', '.join(update_columns)} WHERE id = ?"
                    update_values.append(existing_part[0])
                    cursor.execute(update_query, update_values)
                    count += 1
            else:
                # Dodaj nową część
                columns = list(part_data.keys())
                placeholders = ['?'] * len(columns)
                values = [part_data[column] for column in columns]
                
                insert_query = f"INSERT INTO parts ({', '.join(columns)}) VALUES ({', '.join(placeholders)})"
                cursor.execute(insert_query, values)
                count += 1
        
        # Zatwierdzenie zmian
        conn.commit()
        
        return count
    
    except Exception as e:
        conn.rollback()
        logger.error(f"Błąd podczas importu części: {e}")
        raise


def import_appointments(conn, data_rows, headers):
    """
    Importuje dane wizyt do bazy danych.
    
    Args:
        conn: Połączenie z bazą danych
        data_rows (list): Lista słowników z danymi wizyt
        headers (list): Lista nagłówków
        
    Returns:
        int: Liczba zaimportowanych rekordów
    """
    try:
        cursor = conn.cursor()
        count = 0
        
        # Mapowanie nagłówków do kolumn w bazie danych
        header_mapping = {
            'Klient': 'client_name',
            'Data': 'appointment_date',
            'Godzina': 'appointment_time',
            'Usługa': 'service_type',
            'Status': 'status',
            'Uwagi': 'notes',
            'Czas trwania (min)': 'duration',
            # Angielskie wersje nagłówków
            'Client': 'client_name',
            'Date': 'appointment_date',
            'Time': 'appointment_time',
            'Service': 'service_type',
            'Status': 'status',
            'Notes': 'notes',
            'Duration (min)': 'duration'
        }
        
        # Dodaj wizyty do bazy danych
        for row in data_rows:
            # Przygotuj dane wizyty w formacie odpowiednim dla bazy danych
            appointment_data = {}
            for header in row.keys():
                if header in header_mapping and row[header] is not None:
                    field_name = header_mapping[header]
                    value = row[header]
                    
                    # Konwersja formatu daty
                    if field_name == 'appointment_date':
                        try:
                            if isinstance(value, str):
                                # Sprawdź różne formaty daty
                                if '.' in value:  # Format dd.mm.yyyy
                                    date_obj = datetime.strptime(value, "%d.%m.%Y")
                                elif '-' in value and len(value) >= 10:  # Format yyyy-mm-dd
                                    date_obj = datetime.strptime(value[:10], "%Y-%m-%d")
                                else:
                                    # Spróbuj domyślny format Excela
                                    date_obj = datetime.strptime(value, "%Y-%m-%d")
                                
                                value = date_obj.strftime("%Y-%m-%d")
                            elif hasattr(value, 'strftime'):  # Obiekt daty
                                value = value.strftime("%Y-%m-%d")
                        except (ValueError, TypeError):
                            # Jeśli nie udało się przekonwertować, użyj oryginalnej wartości
                            pass
                    
                    # Konwersja czasu trwania
                    elif field_name == 'duration':
                        try:
                            value = int(value)
                        except (ValueError, TypeError):
                            # Jeśli nie udało się przekonwertować, ustaw domyślny czas trwania (60 minut)
                            value = 60
                    
                    appointment_data[field_name] = value
            
            # Jeśli brak wymaganych danych, pomiń wiersz
            if 'client_name' not in appointment_data or not appointment_data['client_name'] or 'appointment_date' not in appointment_data or not appointment_data['appointment_date']:
                continue
            
            # Znajdź ID klienta na podstawie nazwy
            client_id = None
            if 'client_name' in appointment_data:
                cursor.execute("SELECT id FROM clients WHERE name = ?", (appointment_data['client_name'],))
                client_result = cursor.fetchone()
                
                if client_result:
                    client_id = client_result[0]
                else:
                    # Jeśli klient nie istnieje, dodaj go
                    cursor.execute("INSERT INTO clients (name) VALUES (?)", (appointment_data['client_name'],))
                    client_id = cursor.lastrowid
            
            # Usuń nazwę klienta z danych wizyty
            if 'client_name' in appointment_data:
                del appointment_data['client_name']
            
            # Dodaj ID klienta do danych wizyty
            if client_id:
                appointment_data['client_id'] = client_id
            
            # Sprawdź, czy wizyta już istnieje (po dacie, godzinie i kliencie)
            cursor.execute(
                "SELECT id FROM appointments WHERE appointment_date = ? AND appointment_time = ? AND client_id = ?",
                (appointment_data.get('appointment_date'), appointment_data.get('appointment_time'), client_id)
            )
            existing_appointment = cursor.fetchone()
            
            if existing_appointment:
                # Aktualizuj istniejącą wizytę
                update_columns = []
                update_values = []
                
                for column, value in appointment_data.items():
                    if column not in ['appointment_date', 'appointment_time', 'client_id']:  # Nie aktualizujemy kluczy
                        update_columns.append(f"{column} = ?")
                        update_values.append(value)
                
                if update_columns:
                    update_query = f"UPDATE appointments SET {', '.join(update_columns)} WHERE id = ?"
                    update_values.append(existing_appointment[0])
                    cursor.execute(update_query, update_values)
                    count += 1
            else:
                # Dodaj nową wizytę
                columns = list(appointment_data.keys())
                placeholders = ['?'] * len(columns)
                values = [appointment_data[column] for column in columns]
                
                # Dodaj domyślny status, jeśli nie podano
                if 'status' not in appointment_data:
                    columns.append('status')
                    placeholders.append('?')
                    values.append('Zaplanowana')
                
                insert_query = f"INSERT INTO appointments ({', '.join(columns)}) VALUES ({', '.join(placeholders)})"
                cursor.execute(insert_query, values)
                count += 1
        
        # Zatwierdzenie zmian
        conn.commit()
        
        return count
    
    except Exception as e:
        conn.rollback()
        logger.error(f"Błąd podczas importu wizyt: {e}")
        raise